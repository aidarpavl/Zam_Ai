"""XLSX/CSV оқу, тақырып жолын анықтау, дербес деректерді экспресс-тексеру."""
from __future__ import annotations

import csv
import io
import re

from openpyxl import load_workbook

from .analyzers import to_str, xnorm

MAX_CELLS = 250_000

DEMO_ROWS = [["Сынып", "N", "«5»", "«4»", "«3»", "«2»", "Бір «3»"],
             ["5А", 27, 4, 12, 11, 0, 3],
             ["5Б", 26, 2, 8, 15, 1, 6],
             ["5В", 25, 5, 11, 9, 0, 2]]


def _trim(row):
    row = list(row)
    while row and (row[-1] is None or (isinstance(row[-1], str) and row[-1] == "")):
        row.pop()
    return row


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_csv_text(text: str):
    first = text.splitlines()[0] if text else ""
    delim = max(["\t", ";", ","], key=first.count)
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]


def parse_xlsx(data: bytes):
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"XLSX ашылмады: файл бүлінген немесе қолдау көрсетілмейтін форматта болуы мүмкін ({e}).")
    sheets = []
    try:
        for ws in wb.worksheets:
            rows, cells = [], 0
            for row in ws.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ValueError("Парақ тым үлкен (>250 000 ұяшық). Жүктемені бөліктерге бөліңіз.")
                rows.append(_trim(row))
            sheets.append({"name": ws.title, "rows": rows})
    finally:
        wb.close()
    if not sheets:
        raise ValueError("XLSX ішінде оқылатын парақ табылмады.")
    return sheets


def read_file(name: str, data: bytes):
    low = name.lower()
    if low.endswith(".csv"):
        return [{"name": "CSV", "rows": parse_csv_text(_decode(data))}]
    if low.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    raise ValueError("XLSX және CSV қолдау көрсетіледі. Ескі XLS форматын XLSX ретінде сақтаңыз.")


HEADER_KW = ["класс", "предмет", "всего", "учащ", "качество", "сор", "соч", "пропуск", "результат", "исправ", "питание",
             "ошиб", "оцен", "коррекц", "дифференц", "рефлекс", "активность", "тест", "заплан", "проведено", "ктп",
             "ксп", "журнал", "сынып", "пән", "барлығ", "оқушы", "сапа", "бжб", "тжб", "қалдыр", "нәтиже", "жарамды",
             "тамақ", "қате", "баға", "түзет", "саралау", "белсенд", "жоспар", "өткіз"]


def detect_header_row(rows) -> int:
    best, bscore = 0, -1.0
    for i in range(min(len(rows), 20)):
        vals = [xnorm(v) for v in (rows[i] or [])]
        non = sum(1 for v in vals if v)
        if non < 2:
            continue
        hits = sum(sum(1 for k in HEADER_KW if k in v) for v in vals)
        score = non * .35 + hits * 2.5
        if score > bscore:
            best, bscore = i, score
    return best


def unique_headers(row):
    used, out = {}, []
    for i, v in enumerate(row or []):
        t = to_str(v).strip() or f"Бағана {i + 1}"
        k = t.lower()
        used[k] = used.get(k, 0) + 1
        out.append(f"{t} ({used[k]})" if used[k] > 1 else t)
    return out


def table_from_rows(rows, hi: int):
    headers = unique_headers(rows[hi] if hi < len(rows) else [])
    data = [r for r in rows[hi + 1:] if any(to_str(v).strip() for v in r)]
    return {"headers": headers, "rows": data, "header_index": hi}


PRIVACY_HEADER_KEYS = ["иин", "жсн", "фио", "аты жөні", "толық аты", "телефон", "адрес", "мекенжай", "email", "e mail",
                       "диагноз", "пмпк", "учитель", "педагог", "мұғалім"]
RE_IIN = re.compile(r"(^|\D)\d{12}(\D|$)")
RE_EMAIL = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
RE_PHONE = re.compile(r"(?:\+7|\b8)[\s()\-]*\d{3}[\s()\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}")


def scan_privacy(table):
    bad = []
    for h in table["headers"]:
        n = xnorm(h)
        if any(k in n for k in PRIVACY_HEADER_KEYS) and h not in bad:
            bad.append(h)
    iin = email = phone = 0
    for r in table["rows"][:250]:
        for v in r:
            if not isinstance(v, str):
                continue
            iin += bool(RE_IIN.search(v))
            email += bool(RE_EMAIL.search(v))
            phone += bool(RE_PHONE.search(v))
    return {"badHeaders": bad, "iin": iin, "email": email, "phone": phone,
            "hasRisk": len(bad) + iin + email + phone > 0}
